#!/usr/bin/env python3
"""
Vessel Schedule generator
==========================
Builds "Vessel Schedule.html" (interactive dashboard) and
"Vessel Schedule - <date>-SKED.xlsx" (3-sheet Excel export)
from a daily SKED.xls file, an optional previous day's SKED.xls
(for Early/Delay comparison), and a Wharf.xls code lookup file
(or a cached wharf_lookup.json from a previous run).

USAGE
-----
    python generate_vessel_schedule.py --today 9-9-SKED.xls \
        --yesterday 9-8-SKED.xls \
        --wharf Wharf.xls \
        --now "2026-09-09T12:00:00" \
        --outdir .

Only --today is required. If --yesterday is omitted, the Early/Delay
feature will simply show no changes. If --wharf is omitted, the script
looks for a cached wharf_lookup.json next to this script (built by an
earlier run) and reuses it; if neither is available, wharf codes will
just show without a resolved name.

If --now is omitted, the script uses the current system time.

REQUIRES: pandas, xlrd (for legacy .xls), openpyxl
    pip install pandas xlrd openpyxl --break-system-packages
"""

import argparse
import json
import sys
from pathlib import Path
from datetime import datetime

import pandas as pd

SCRIPT_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = SCRIPT_DIR / "vessel_schedule_template.html"
WHARF_CACHE_PATH = SCRIPT_DIR / "wharf_lookup.json"
LOGO_PATH = SCRIPT_DIR / "logo ha.png"


def logo_data_uri():
    """Return a base64 data URI for the header logo, or '' if not present."""
    if not LOGO_PATH.exists():
        print(f"[warn] {LOGO_PATH.name} not found; header logo will be hidden")
        return ""
    import base64
    b64 = base64.b64encode(LOGO_PATH.read_bytes()).decode("ascii")
    return f"data:image/png;base64,{b64}"


# ---------------------------------------------------------------------------
# Data loading helpers
# ---------------------------------------------------------------------------

def load_sked(path):
    df = pd.read_excel(path, sheet_name=0, header=0)
    for c in ["ETA Date", "ETB Date", "ETD Date"]:
        df[c] = pd.to_datetime(df[c])
    return df


def load_wharf_lookup(wharf_path):
    """Build {code: {name, area, country}} from Wharf.xls, or fall back
    to a cached JSON from a previous run."""
    if wharf_path:
        df = pd.read_excel(wharf_path, sheet_name=0, header=0)
        df.columns = ["Wharf", "WharfName", "Area", "AreaName", "Country", "CountryName"]
        df = df.dropna(subset=["Wharf"])
        lookup = {}
        for _, r in df.iterrows():
            lookup[r["Wharf"]] = {
                "name": None if pd.isna(r["WharfName"]) else str(r["WharfName"]).strip(),
                "area": None if pd.isna(r["AreaName"]) else str(r["AreaName"]).strip(),
                "country": None if pd.isna(r["CountryName"]) else str(r["CountryName"]).strip(),
            }
        # refresh the cache for next time
        with open(WHARF_CACHE_PATH, "w", encoding="utf-8") as f:
            json.dump(lookup, f, ensure_ascii=False)
        return lookup

    if WHARF_CACHE_PATH.exists():
        print(f"[info] --wharf not given, reusing cached {WHARF_CACHE_PATH.name}")
        with open(WHARF_CACHE_PATH, encoding="utf-8") as f:
            return json.load(f)

    print("[warn] no wharf lookup available; wharf codes will show unresolved")
    return {}


# ---------------------------------------------------------------------------
# Core transforms (mirrors the dashboard's own JS logic)
# ---------------------------------------------------------------------------

def build_vessel_summary(df, now):
    rows = []
    for vessel, g in df.groupby("Vessel Name"):
        if vessel == "TO BE NOMINATED":
            continue
        g = g.sort_values("ETA Date").reset_index(drop=True)

        docked = g[(g["ETA Date"] <= now) & (g["ETD Date"] >= now)]
        status = "unknown"
        loc = wharf = since_eta = until_etd = dest_pod = docked_vyg = None
        transit_from = transit_to = transit_eta = None

        if len(docked) > 0:
            d = docked.iloc[-1]
            status = "docked"
            loc, wharf = d["POL"], d["Wharf"]
            since_eta, until_etd, dest_pod = d["ETA Date"], d["ETD Date"], d["POD"]
            docked_vyg = d["Vyg Bound"]
        else:
            past = g[g["ETD Date"] <= now]
            fut = g[g["ETA Date"] >= now]
            if len(past) > 0 and len(fut) > 0:
                a, b = past.iloc[-1], fut.iloc[0]
                status = "transit"
                transit_from, transit_to, transit_eta = a["POL"], b["POL"], b["ETA Date"]
            elif len(past) > 0:
                a = past.iloc[-1]
                status = "transit"
                transit_from, transit_to = a["POL"], a["POD"]
            else:
                status = "nodata"

        fut2 = g[g["ETA Date"] > now]
        bkk = fut2[fut2["POL"] == "THBKK"].sort_values("ETA Date")
        lch = fut2[fut2["POL"] == "THLCH"].sort_values("ETA Date")
        bkk_row = bkk.iloc[0] if len(bkk) else None
        lch_row = lch.iloc[0] if len(lch) else None

        code = g.iloc[0]["Vessel"]
        rows.append({
            "vessel": vessel,
            "vesselCode": None if pd.isna(code) else str(code).strip(),
            "service": g.iloc[0]["Service"],
            "opLiner": g.iloc[0]["Op.Liner"],
            "status": status,
            "dockedLoc": loc, "dockedWharf": wharf,
            "dockedSince": since_eta.isoformat() if since_eta is not None else None,
            "dockedUntil": until_etd.isoformat() if until_etd is not None else None,
            "dockedNextDest": dest_pod,
            "dockedVygBound": None if docked_vyg is None or pd.isna(docked_vyg) else str(docked_vyg),
            "transitFrom": transit_from, "transitTo": transit_to,
            "transitEta": transit_eta.isoformat() if transit_eta is not None else None,
            "bkkEta": bkk_row["ETA Date"].isoformat() if bkk_row is not None else None,
            "bkkWharf": bkk_row["Wharf"] if bkk_row is not None else None,
            "bkkVygBound": str(bkk_row["Vyg Bound"]) if bkk_row is not None and not pd.isna(bkk_row["Vyg Bound"]) else None,
            "lchEta": lch_row["ETA Date"].isoformat() if lch_row is not None else None,
            "lchWharf": lch_row["Wharf"] if lch_row is not None else None,
            "lchVygBound": str(lch_row["Vyg Bound"]) if lch_row is not None and not pd.isna(lch_row["Vyg Bound"]) else None,
        })
    return rows


def build_legs(df):
    legs = {}
    for vessel, g in df.groupby("Vessel Name"):
        if vessel == "TO BE NOMINATED":
            continue
        g = g.sort_values("ETA Date")
        rows = []
        for _, r in g.iterrows():
            rows.append({
                "service": r["Service"], "vyg": str(r["Vyg"]), "bound": r["Bound"],
                "vygBound": r["Vyg Bound"], "wharf": r["Wharf"],
                "pol": r["POL"], "pod": r["POD"],
                "eta": r["ETA Date"].isoformat() if not pd.isna(r["ETA Date"]) else None,
                "etb": r["ETB Date"].isoformat() if not pd.isna(r["ETB Date"]) else None,
                "etd": r["ETD Date"].isoformat() if not pd.isna(r["ETD Date"]) else None,
            })
        legs[vessel] = rows
    return legs


def build_schedule_changes(df_old, df_new):
    old_map = {}
    for _, r in df_old.iterrows():
        if pd.isna(r["Vyg Bound"]) or pd.isna(r["POL"]) or pd.isna(r["Vessel Name"]):
            continue
        key = f"{r['Vessel Name']}|{r['Vyg Bound']}|{r['POL']}"
        old_map[key] = r["ETA Date"]

    changes = {}
    for _, r in df_new.iterrows():
        if pd.isna(r["Vyg Bound"]) or pd.isna(r["POL"]) or pd.isna(r["Vessel Name"]):
            continue
        key = f"{r['Vessel Name']}|{r['Vyg Bound']}|{r['POL']}"
        if key in old_map:
            old_eta, new_eta = old_map[key], r["ETA Date"]
            if pd.isna(old_eta) or pd.isna(new_eta):
                continue
            delta_hours = (new_eta - old_eta).total_seconds() / 3600
            if abs(delta_hours) >= 0.5:
                entry = {
                    "pol": r["POL"], "pod": r["POD"], "vygBound": r["Vyg Bound"],
                    "oldEta": old_eta.isoformat(), "newEta": new_eta.isoformat(),
                    "deltaHours": round(delta_hours, 1),
                }
                changes.setdefault(r["Vessel Name"], []).append(entry)

    for v in changes:
        changes[v].sort(key=lambda e: e["newEta"])
    return changes


# ---------------------------------------------------------------------------
# HTML dashboard output
# ---------------------------------------------------------------------------

def write_dashboard(vessels, legs, changes, wharf_lookup, now, today_name, yesterday_name, outdir):
    if not TEMPLATE_PATH.exists():
        print(f"[error] template not found at {TEMPLATE_PATH}. Keep "
              f"vessel_schedule_template.html next to this script.", file=sys.stderr)
        sys.exit(1)

    with open(TEMPLATE_PATH, encoding="utf-8") as f:
        html = f.read()

    if yesterday_name:
        source_label = f"{today_name} (vs {yesterday_name}) · Wharf.xls"
    else:
        source_label = f"{today_name} · Wharf.xls"

    html = html.replace("__LOGO_DATA_URI__", logo_data_uri())
    html = html.replace("__NOW_ISO__", now.isoformat())
    html = html.replace("__SCHEDULE_CHANGES_JSON__", json.dumps(changes))
    html = html.replace("__LEGS_JSON__", json.dumps(legs))
    html = html.replace("__VESSELS_JSON__", json.dumps(vessels))
    html = html.replace("__WHARF_LOOKUP_JSON__", json.dumps(wharf_lookup, ensure_ascii=False))
    html = html.replace("__SOURCE_LABEL__", source_label)
    html = html.replace("__SOURCE_FILENAME__", today_name)

    out_path = Path(outdir) / "Vessel Schedule.html"
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)
    return out_path


# ---------------------------------------------------------------------------
# Excel output (Summary / Full Schedule / Early-Delay)
# ---------------------------------------------------------------------------

def write_excel(vessels, legs, changes, today_name, yesterday_name, outdir):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    def parse_dt(s):
        if not s:
            return None
        try:
            return datetime.fromisoformat(s)
        except Exception:
            return s

    summary_rows = []
    for v in vessels:
        if v["status"] == "docked":
            loc = f"{v['dockedLoc']} ({v['dockedWharf']})"
            note = f"Alongside, departs to {v['dockedNextDest']} ~{v['dockedUntil']}"
        elif v["status"] == "transit":
            loc = f"{v['transitFrom']} -> {v['transitTo']}"
            note = f"In transit, ETA {v['transitEta']}" if v["transitEta"] else "In transit"
        else:
            loc, note = "-", "No data in window"
        summary_rows.append({
            "Vessel Name": v["vessel"], "Code": v.get("vesselCode"),
            "Service": v["service"], "Op.Liner": v["opLiner"],
            "Status": v["status"], "Current Location": loc,
            "Vyg Bound": v.get("dockedVygBound"), "Note": note,
            "ETA THBKK": parse_dt(v["bkkEta"]), "Wharf THBKK": v["bkkWharf"],
            "Vyg Bound THBKK": v.get("bkkVygBound"),
            "ETA THLCH": parse_dt(v["lchEta"]), "Wharf THLCH": v["lchWharf"],
            "Vyg Bound THLCH": v.get("lchVygBound"),
        })
    df_summary = pd.DataFrame(summary_rows)

    leg_rows = []
    for vessel, ls in legs.items():
        for l in ls:
            leg_rows.append({
                "Vessel Name": vessel, "Service": l["service"], "Voyage": l["vygBound"],
                "Wharf": l["wharf"], "POL": l["pol"], "POD": l["pod"],
                "ETA": parse_dt(l["eta"]), "ETB": parse_dt(l["etb"]), "ETD": parse_dt(l["etd"]),
            })
    df_legs = pd.DataFrame(leg_rows).sort_values(["Vessel Name", "ETA"]) if leg_rows else pd.DataFrame(
        columns=["Vessel Name", "Service", "Voyage", "Wharf", "POL", "POD", "ETA", "ETB", "ETD"])

    old_col = f"Old ETA ({yesterday_name})" if yesterday_name else "Old ETA"
    new_col = f"New ETA ({today_name})"
    change_rows = []
    for vessel, entries in changes.items():
        for e in entries:
            change_rows.append({
                "Vessel Name": vessel, "POL": e["pol"], "POD": e["pod"], "Voyage": e["vygBound"],
                old_col: parse_dt(e["oldEta"]), new_col: parse_dt(e["newEta"]),
                "Delta (hours)": e["deltaHours"], "Status": "Delay" if e["deltaHours"] > 0 else "Early",
            })
    df_changes = pd.DataFrame(change_rows).sort_values(["Vessel Name", new_col]) if change_rows else pd.DataFrame(
        columns=["Vessel Name", "POL", "POD", "Voyage", old_col, new_col, "Delta (hours)", "Status"])

    wb = Workbook()
    wb.remove(wb.active)
    HEADER_FILL = PatternFill(start_color="1C3A54", end_color="1C3A54", fill_type="solid")
    HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    BODY_FONT = Font(name="Arial", size=10)
    DATE_FMT = "DD-MMM HH:MM"

    def write_sheet(name, df, date_cols):
        ws = wb.create_sheet(name)
        for j, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=j, value=col)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for i, row in enumerate(df.itertuples(index=False), start=2):
            for j, val in enumerate(row, start=1):
                cell = ws.cell(row=i, column=j, value=val)
                cell.font = BODY_FONT
                if df.columns[j - 1] in date_cols and val is not None:
                    cell.number_format = DATE_FMT
        for j, col in enumerate(df.columns, start=1):
            maxlen = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).tolist()[:200]] or [10])
            ws.column_dimensions[get_column_letter(j)].width = min(max(maxlen + 2, 10), 42)
        ws.freeze_panes = "A2"

    write_sheet("Summary", df_summary, {"ETA THBKK", "ETA THLCH"})
    write_sheet("Full Schedule", df_legs, {"ETA", "ETB", "ETD"})
    write_sheet("Early-Delay", df_changes, {old_col, new_col})

    date_tag = today_name.replace(".xls", "").replace(".xlsx", "")
    out_path = Path(outdir) / f"Vessel Schedule - {date_tag}.xlsx"
    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--today", required=True, help="Today's SKED.xls file")
    ap.add_argument("--yesterday", help="Previous day's SKED.xls, for Early/Delay comparison")
    ap.add_argument("--wharf", help="Wharf.xls code lookup file (omit to reuse cached wharf_lookup.json)")
    ap.add_argument("--now", help="Reference timestamp, ISO format e.g. 2026-09-09T12:00:00 (default: current time)")
    ap.add_argument("--outdir", default=".", help="Output directory (default: current directory)")
    args = ap.parse_args()

    today_path = Path(args.today)
    now = datetime.fromisoformat(args.now) if args.now else datetime.now()

    print(f"[1/5] Loading {today_path.name} ...")
    df_today = load_sked(today_path)

    df_yesterday = None
    if args.yesterday:
        print(f"[2/5] Loading {Path(args.yesterday).name} for comparison ...")
        df_yesterday = load_sked(args.yesterday)
    else:
        print("[2/5] No --yesterday given, skipping Early/Delay comparison")

    print("[3/5] Loading wharf lookup ...")
    wharf_lookup = load_wharf_lookup(args.wharf)

    print("[4/5] Building vessel summary, legs, and schedule changes ...")
    vessels = build_vessel_summary(df_today, now)
    legs = build_legs(df_today)
    changes = build_schedule_changes(df_yesterday, df_today) if df_yesterday is not None else {}

    missing_wharves = sorted({w for w in df_today["Wharf"].dropna().unique() if w not in wharf_lookup})
    if missing_wharves:
        print(f"[warn] {len(missing_wharves)} wharf code(s) not found in lookup: {missing_wharves}")

    print("[5/5] Writing outputs ...")
    yesterday_name = Path(args.yesterday).name if args.yesterday else None
    html_path = write_dashboard(vessels, legs, changes, wharf_lookup, now,
                                 today_path.name, yesterday_name, args.outdir)
    xlsx_path = write_excel(vessels, legs, changes, today_path.name, yesterday_name, args.outdir)

    print()
    print(f"Done: {len(vessels)} vessels, {sum(len(v) for v in legs.values())} legs, "
          f"{len(changes)} vessel(s) with schedule changes.")
    print(f"  -> {html_path}")
    print(f"  -> {xlsx_path}")


if __name__ == "__main__":
    main()
